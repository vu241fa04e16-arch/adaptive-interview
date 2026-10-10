import hashlib
import io
import os
import re
from pathlib import Path
from typing import Tuple
from fastapi import UploadFile
import pypdf
import docx

from app.config import get_settings
from app.core.exceptions import (
    FileSizeLimitExceededException,
    InvalidFileException,
)
from app.core.logging import logger
from app.core.ollama_client import AsyncOllamaClient
from app.schemas.resume import ResumeAnalysisLLMOutput

settings = get_settings()


class ResumeService:
    """Service for validating, parsing, and extracting technical competencies from candidate resumes."""

    def __init__(self, ollama_client: AsyncOllamaClient):
        self.ollama = ollama_client

    async def validate_and_read_upload(self, file: UploadFile) -> Tuple[bytes, str, str]:
        """Validates file extension, size, and computes SHA-256 hash.
        
        Returns:
            Tuple of (raw_bytes, sanitized_filename, sha256_hash)
        """
        filename = file.filename or "uploaded_resume.pdf"
        # Sanitize filename to prevent directory traversal
        clean_filename = os.path.basename(filename)
        clean_filename = re.sub(r"[^\w\.-]", "_", clean_filename)

        ext = Path(clean_filename).suffix.lower()
        if ext not in settings.ALLOWED_RESUME_EXTENSIONS:
            raise InvalidFileException(
                message=f"Unsupported file extension '{ext}'. Only PDF (.pdf) and DOCX (.docx) files are supported.",
                status_code=415,
                error_code="UNSUPPORTED_MEDIA_TYPE",
                details={"allowed_extensions": settings.ALLOWED_RESUME_EXTENSIONS, "received_filename": clean_filename},
            )

        content = await file.read()
        file_size = len(content)

        if file_size == 0:
            raise InvalidFileException(
                message="Uploaded file is empty (0 bytes).",
                status_code=400,
                error_code="EMPTY_FILE",
            )

        max_size_bytes = settings.MAX_RESUME_SIZE_MB * 1024 * 1024
        if file_size > max_size_bytes:
            raise FileSizeLimitExceededException(
                max_size_mb=settings.MAX_RESUME_SIZE_MB,
                actual_size_bytes=file_size,
            )

        file_hash = hashlib.sha256(content).hexdigest()
        return content, clean_filename, file_hash

    def extract_text(self, content: bytes, filename: str) -> str:
        """Extracts text content from PDF or DOCX file bytes."""
        ext = Path(filename).suffix.lower()
        extracted_text = ""

        try:
            if ext == ".pdf":
                pdf_stream = io.BytesIO(content)
                reader = pypdf.PdfReader(pdf_stream)
                if reader.is_encrypted:
                    try:
                        reader.decrypt("")
                    except Exception as e:
                        raise InvalidFileException(
                            message="PDF document is password-protected and cannot be analyzed.",
                            status_code=400,
                            error_code="PASSWORD_PROTECTED_DOCUMENT",
                        ) from e

                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_text += page_text + "\n"

            elif ext == ".docx":
                docx_stream = io.BytesIO(content)
                doc = docx.Document(docx_stream)
                for paragraph in doc.paragraphs:
                    if paragraph.text:
                        extracted_text += paragraph.text + "\n"
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            extracted_text += row_text + "\n"

        except InvalidFileException:
            raise
        except Exception as e:
            logger.error(f"Error parsing document {filename}: {e}")
            raise InvalidFileException(
                message=f"Failed to parse document: the file may be corrupted or malformed ({str(e)}).",
                status_code=422,
                error_code="DOCUMENT_PARSE_ERROR",
                details={"filename": filename, "error": str(e)},
            ) from e

        cleaned_text = extracted_text.strip()
        if len(cleaned_text) < 30:
            raise InvalidFileException(
                message="Document contains insufficient readable text (fewer than 30 characters). Ensure it is not an image-only scan.",
                status_code=422,
                error_code="INSUFFICIENT_TEXT_CONTENT",
                details={"extracted_length": len(cleaned_text)},
            )

        return cleaned_text

    def save_upload_safely(self, session_id: str, filename: str, file_hash: str, content: bytes) -> str:
        """Saves uploaded resume to isolated secure directory."""
        upload_base = Path(settings.UPLOAD_DIR) / "resumes" / session_id
        upload_base.mkdir(parents=True, exist_ok=True)

        safe_name = f"{file_hash[:12]}_{filename}"
        destination = upload_base / safe_name

        destination.write_bytes(content)
        return str(destination)

    async def analyze_resume_with_ollama(
        self,
        resume_text: str,
        target_role: str,
        target_level: str,
    ) -> ResumeAnalysisLLMOutput:
        """Invokes real Ollama local service to extract claimed skills, projects, and suggested focus areas."""
        system_prompt = (
            "You are an expert technical interviewer and talent architect.\n"
            "Analyze the candidate's resume text and extract ONLY technical information explicitly stated in the document.\n\n"
            "CRITICAL RULES:\n"
            "1. NEVER fabricate or assume skills, projects, or experience years. If a skill is not mentioned, omit it.\n"
            "2. If years of experience for a skill are not explicitly stated, set claimed_experience_years to null.\n"
            "3. Every claimed skill must have a verbatim context_snippet quote from the resume proving its presence.\n"
            "4. Suggest 3 to 6 high-value focus areas that should be deeply verified or probed during the interview "
            f"for the target role '{target_role}' at '{target_level}' seniority.\n"
            "5. Output MUST adhere strictly to the JSON Schema provided."
        )

        prompt = (
            f"TARGET ROLE: {target_role}\n"
            f"TARGET SENIORITY LEVEL: {target_level}\n\n"
            "CANDIDATE RESUME TEXT:\n"
            "------------------------\n"
            f"{resume_text}\n"
            "------------------------\n\n"
            "Extract candidate overview, all claimed technical skills with supporting snippets, "
            "notable projects, and targeted interview focus areas."
        )

        analysis: ResumeAnalysisLLMOutput = await self.ollama.generate_structured(
            prompt=prompt,
            schema=ResumeAnalysisLLMOutput,
            system_prompt=system_prompt,
            options={"temperature": 0.1, "num_predict": 1024},
        )

        return analysis
